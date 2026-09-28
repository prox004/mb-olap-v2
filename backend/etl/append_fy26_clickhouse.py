"""
MB-OLAP V2: Incremental Append ETL Pipeline for ClickHouse
Appends FY26 Sales and Stock Movement records to existing ClickHouse database without touching historical FY25 data.
"""

import os
import sys
import time
import datetime
from decimal import Decimal
import openpyxl
import pandas as pd
import clickhouse_connect

# Ensure UTF-8 stdout encoding on Windows
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

# Add project root to sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

MONTH_DATE_MAP_FY26 = {
    "Apr Q2-26": (datetime.date(2026, 4, 1), datetime.date(2026, 4, 30)),
    "May Q2-26": (datetime.date(2026, 5, 1), datetime.date(2026, 5, 31)),
    "Jun Q2-26": (datetime.date(2026, 6, 1), datetime.date(2026, 6, 30)),
    "Jul Q3-26": (datetime.date(2026, 7, 1), datetime.date(2026, 7, 31)),
    "Aug Q3-26": (datetime.date(2026, 8, 1), datetime.date(2026, 8, 31)),
    "Sep Q3-26": (datetime.date(2026, 9, 1), datetime.date(2026, 9, 15)),
}

DATES_FY26 = [
    {
        "date": datetime.date(2026, 4, 1),
        "year": 2026,
        "quarter": 2,
        "month": 4,
        "month_name": "April",
        "month_period_label": "Apr Q2-26",
        "day_of_month": 1,
        "day_of_week": 3,
        "day_name": "Wednesday",
        "is_weekend": 0,
    },
    {
        "date": datetime.date(2026, 5, 1),
        "year": 2026,
        "quarter": 2,
        "month": 5,
        "month_name": "May",
        "month_period_label": "May Q2-26",
        "day_of_month": 1,
        "day_of_week": 5,
        "day_name": "Friday",
        "is_weekend": 0,
    },
    {
        "date": datetime.date(2026, 6, 1),
        "year": 2026,
        "quarter": 2,
        "month": 6,
        "month_name": "June",
        "month_period_label": "Jun Q2-26",
        "day_of_month": 1,
        "day_of_week": 1,
        "day_name": "Monday",
        "is_weekend": 0,
    },
    {
        "date": datetime.date(2026, 7, 1),
        "year": 2026,
        "quarter": 3,
        "month": 7,
        "month_name": "July",
        "month_period_label": "Jul Q3-26",
        "day_of_month": 1,
        "day_of_week": 3,
        "day_name": "Wednesday",
        "is_weekend": 0,
    },
    {
        "date": datetime.date(2026, 8, 1),
        "year": 2026,
        "quarter": 3,
        "month": 8,
        "month_name": "August",
        "month_period_label": "Aug Q3-26",
        "day_of_month": 1,
        "day_of_week": 6,
        "day_name": "Saturday",
        "is_weekend": 1,
    },
    {
        "date": datetime.date(2026, 9, 1),
        "year": 2026,
        "quarter": 3,
        "month": 9,
        "month_name": "September",
        "month_period_label": "Sep Q3-26",
        "day_of_month": 1,
        "day_of_week": 2,
        "day_name": "Tuesday",
        "is_weekend": 0,
    },
]


def get_clickhouse_client():
    from backend.app.config import settings
    host = os.getenv("CLICKHOUSE_HOST", getattr(settings, "CLICKHOUSE_HOST", "localhost"))
    port = int(os.getenv("CLICKHOUSE_PORT", getattr(settings, "CLICKHOUSE_PORT", 8123)))
    database = os.getenv("CLICKHOUSE_DATABASE", getattr(settings, "CLICKHOUSE_DATABASE", "mb_olap_v2"))
    user = os.getenv("CLICKHOUSE_USER", getattr(settings, "CLICKHOUSE_USER", "default"))
    password = os.getenv("CLICKHOUSE_PASSWORD", getattr(settings, "CLICKHOUSE_PASSWORD", ""))
    
    print(f"Connecting to ClickHouse at {host}:{port} (DB: {database})...")
    client = clickhouse_connect.get_client(
        host=host,
        port=port,
        database=database,
        username=user,
        password=password,
        connect_timeout=15,
        send_receive_timeout=300,
    )
    return client


def append_dim_date(client):
    print("\n--- 1. Appending dim_date for FY26 ---")
    existing_dates = set(r[0] for r in client.query("SELECT date FROM dim_date").result_rows)
    new_dates = [d for d in DATES_FY26 if d["date"] not in existing_dates]
    if new_dates:
        df_new_dates = pd.DataFrame(new_dates)
        client.insert("dim_date", df_new_dates, column_names=list(df_new_dates.columns))
        print(f"   -> Inserted {len(new_dates)} new date records for FY26.")
    else:
        print("   -> FY26 dates already present in dim_date.")
    total_dates = client.command("SELECT count(*) FROM dim_date")
    print(f"   Total dim_date count: {total_dates}")


def append_sales_and_products(client, sales_file: str):
    print(f"\n--- 2. Reading FY26 Sales Data from {sales_file} ---")
    sheet_name = "1-APR TO 15-SEP (26) - detail"
    wb = openpyxl.load_workbook(sales_file, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {sales_file}. Found: {wb.sheetnames}")

    sheet = wb[sheet_name]
    t0 = time.time()
    
    fact_rows = []
    items_dict = {}
    valid_count = 0
    total_revenue = 0.0
    total_qty = 0

    print("   Streaming rows from Excel...")
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if idx < 6:
            continue
        if idx == 6:
            # Header row
            continue

        # Check for empty or summary termination
        if row[0] is None and row[1] is None:
            print(f"   Termination/summary row reached at row {idx}")
            break

        store_code = str(row[1]).strip() if row[1] else ""
        bill_qty = int(row[2]) if row[2] is not None else 0
        net_amt = float(row[3]) if row[3] is not None else 0.0
        cogs = float(row[4]) if row[4] is not None else 0.0
        division = str(row[5]).strip() if row[5] else ""
        section = str(row[6]).strip() if row[6] else ""
        department = str(row[7]).strip() if row[7] else ""
        group_alias = str(row[8]).strip() if row[8] else ""
        article_name = str(row[9]).strip() if row[9] else ""
        item_code = str(row[10]).strip() if row[10] else ""
        cat1 = str(row[11]).strip() if row[11] else ""
        cat2 = str(row[12]).strip() if row[12] else ""
        cat3 = str(row[13]).strip() if row[13] else None
        cat4 = str(row[14]).strip() if row[14] else None
        cat5 = str(row[15]).strip() if row[15] else None
        cat6 = str(row[16]).strip() if row[16] else ""
        rsp = float(row[17]) if row[17] is not None else 0.0
        desc1 = str(row[18]).strip() if row[18] else ""
        desc2 = str(row[19]).strip() if row[19] else None
        desc3 = str(row[20]).strip() if row[20] else None

        gen_val = row[21]
        gen_date = gen_val.date() if isinstance(gen_val, (datetime.datetime, datetime.date)) else None

        stock_val = row[22]
        stock_date = stock_val.date() if isinstance(stock_val, (datetime.datetime, datetime.date)) else None

        month_label = str(row[23]).strip() if len(row) > 23 and row[23] else ""
        if month_label not in MONTH_DATE_MAP_FY26:
            continue

        start_date, end_date = MONTH_DATE_MAP_FY26[month_label]
        gross_profit = round(net_amt - cogs, 2)

        # Collect product dimension
        if item_code and item_code not in items_dict:
            items_dict[item_code] = {
                "item_code": item_code,
                "article_name": article_name,
                "division": division,
                "section": section,
                "department": department,
                "group_alias": group_alias,
                "category1": cat1,
                "category2": cat2,
                "category3": cat3,
                "category4": cat4,
                "category5": cat5,
                "category6": cat6,
                "colour": desc1 if desc1 else "NA",
                "vendor_name": desc2,
                "style_code": desc3,
                "rsp": rsp,
                "generated_date": gen_date,
                "last_stock_in_date": stock_date,
            }

        # Collect fact row
        fact_rows.append({
            "period_start_date": start_date,
            "period_end_date": end_date,
            "period_month_label": month_label,
            "store_code": store_code,
            "item_code": item_code,
            "bill_qty": bill_qty,
            "net_amount": net_amt,
            "cogs": cogs,
            "gross_profit": gross_profit,
            "unit_rsp": rsp,
        })

        valid_count += 1
        total_revenue += net_amt
        total_qty += bill_qty

        if valid_count % 100000 == 0:
            print(f"   Read {valid_count:,} FY26 sales rows ({time.time() - t0:.1f}s)...")

    elapsed_read = time.time() - t0
    print(f"   -> Finished reading {valid_count:,} FY26 sales rows in {elapsed_read:.2f}s")
    print(f"   -> Found {len(items_dict):,} distinct FY26 product items")
    print(f"   -> Total FY26 Net Revenue: ₹{total_revenue:,.2f}, Total Units: {total_qty:,}")

    # 1. Insert Products into dim_product
    print("\n--- 3. Updating dim_product with FY26 items ---")
    df_items = pd.DataFrame(list(items_dict.values()))
    df_items["rsp"] = df_items["rsp"].astype(float)
    batch_size = 50000
    for i in range(0, len(df_items), batch_size):
        chunk = df_items.iloc[i : i + batch_size]
        client.insert("dim_product", chunk, column_names=list(chunk.columns))
    total_products = client.command("SELECT count(*) FROM dim_product")
    print(f"   -> dim_product now contains {total_products:,} total items (ReplacingMergeTree).")

    # 2. Append FY26 Sales to fact_sales_monthly
    print("\n--- 4. Appending FY26 sales to fact_sales_monthly ---")
    # Idempotent cleanup: if FY26 partitions (202604 - 202609) already exist, drop those specific partitions
    existing_partitions = client.query("SELECT DISTINCT partition FROM system.parts WHERE table = 'fact_sales_monthly' AND active = 1").result_rows
    for part in ["202604", "202605", "202606", "202607", "202608", "202609"]:
        if any(r[0] == part for r in existing_partitions):
            print(f"   Dropping existing partition {part} for clean reload...")
            client.command(f"ALTER TABLE fact_sales_monthly DROP PARTITION '{part}'")

    df_fact = pd.DataFrame(fact_rows)
    df_fact["bill_qty"] = df_fact["bill_qty"].astype(int)
    df_fact["net_amount"] = df_fact["net_amount"].astype(float)
    df_fact["cogs"] = df_fact["cogs"].astype(float)
    df_fact["gross_profit"] = df_fact["gross_profit"].astype(float)
    df_fact["unit_rsp"] = df_fact["unit_rsp"].astype(float)

    for i in range(0, len(df_fact), batch_size):
        chunk = df_fact.iloc[i : i + batch_size]
        client.insert("fact_sales_monthly", chunk, column_names=list(chunk.columns))
        print(f"   Inserted sales rows {i+1:,} to {min(i+batch_size, len(df_fact)):,}...")

    total_sales_rows = client.command("SELECT count(*) FROM fact_sales_monthly")
    print(f"   -> fact_sales_monthly now has {total_sales_rows:,} total rows (FY25 + FY26)!")


def append_inventory(client, inv_file: str):
    print(f"\n--- 5. Reading FY26 Inventory Data from {inv_file} ---")
    sheet_name = "26"
    wb = openpyxl.load_workbook(inv_file, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {inv_file}. Found: {wb.sheetnames}")

    sheet = wb[sheet_name]
    t0 = time.time()

    inv_rows = []
    start_date = datetime.date(2026, 4, 1)
    end_date = datetime.date(2026, 9, 15)
    report_label = "2026-04-01 to 2026-09-15"

    print("   Streaming rows from Excel sheet '26'...")
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if idx < 8:
            continue

        item_code_val = row[0]
        if item_code_val is None:
            continue
        item_code = str(item_code_val).strip()
        if item_code.lower() == "sum":
            print(f"   Summary total row reached at row {idx}")
            break

        store_code = str(row[1]).strip() if row[1] else ""
        if not store_code:
            continue

        opening_qty = int(row[2]) if row[2] is not None else 0
        opening_amt = float(row[3]) if row[3] is not None else 0.0
        purchase_net_qty = int(row[4]) if row[4] is not None else 0
        purchase_net_amt = float(row[5]) if row[5] is not None else 0.0
        transfer_in_qty = int(row[6]) if row[6] is not None else 0
        transfer_in_amt = float(row[7]) if row[7] is not None else 0.0
        transfer_out_qty = int(row[8]) if row[8] is not None else 0
        transfer_out_amt = float(row[9]) if row[9] is not None else 0.0
        cogca_qty = int(row[10]) if row[10] is not None else 0
        cogca_amt = float(row[11]) if row[11] is not None else 0.0
        closing_qty = int(row[12]) if row[12] is not None else 0
        final_sale_qty = int(row[13]) if row[13] is not None else 0
        closing_amt = float(row[14]) if row[14] is not None else 0.0
        transit_qty = int(row[15]) if len(row) > 15 and row[15] is not None else 0
        transit_amt = float(row[16]) if len(row) > 16 and row[16] is not None else 0.0

        inv_rows.append({
            "period_start_date": start_date,
            "period_end_date": end_date,
            "report_period_label": report_label,
            "store_code": store_code,
            "item_code": item_code,
            "opening_qty": opening_qty,
            "opening_amt": opening_amt,
            "purchase_net_qty": purchase_net_qty,
            "purchase_net_amt": purchase_net_amt,
            "transfer_in_qty": transfer_in_qty,
            "transfer_in_amt": transfer_in_amt,
            "transfer_out_qty": transfer_out_qty,
            "transfer_out_amt": transfer_out_amt,
            "cogca_qty": cogca_qty,
            "cogca_amt": cogca_amt,
            "closing_qty": closing_qty,
            "final_sale_qty": final_sale_qty,
            "closing_amt": closing_amt,
            "transit_qty": transit_qty,
            "transit_amt": transit_amt,
        })

        if len(inv_rows) % 100000 == 0:
            print(f"   Read {len(inv_rows):,} FY26 inventory rows ({time.time() - t0:.1f}s)...")

    elapsed_read = time.time() - t0
    print(f"   -> Finished reading {len(inv_rows):,} FY26 inventory rows in {elapsed_read:.2f}s")

    # Appending to fact_inventory
    print("\n--- 6. Appending FY26 inventory to fact_inventory ---")
    # Idempotent cleanup: delete existing FY26 inventory if present
    existing_fy26_count = client.command("SELECT count(*) FROM fact_inventory WHERE period_start_date = '2026-04-01'")
    if int(existing_fy26_count) > 0:
        print(f"   Removing existing FY26 inventory ({existing_fy26_count:,} rows) for clean reload...")
        client.command("ALTER TABLE fact_inventory DELETE WHERE period_start_date = '2026-04-01'")
        time.sleep(1)

    df_inv = pd.DataFrame(inv_rows)
    batch_size = 50000
    for i in range(0, len(df_inv), batch_size):
        chunk = df_inv.iloc[i : i + batch_size]
        client.insert("fact_inventory", chunk, column_names=list(chunk.columns))
        print(f"   Inserted inventory rows {i+1:,} to {min(i+batch_size, len(df_inv)):,}...")

    total_inv_rows = client.command("SELECT count(*) FROM fact_inventory")
    print(f"   -> fact_inventory now has {total_inv_rows:,} total rows (FY25 + FY26)!")


def verify_clickhouse(client):
    print("\n==================================================")
    print("CLICKHOUSE VERIFICATION AFTER INCREMENTAL APPEND")
    print("==================================================")

    periods_sales = client.query("""
        SELECT toYear(period_start_date) as yr, period_month_label, count(), round(sum(net_amount), 2), sum(bill_qty)
        FROM fact_sales_monthly
        GROUP BY yr, period_month_label, period_start_date
        ORDER BY period_start_date
    """).result_rows

    print("\nSales By Month Breakdown:")
    for yr, label, cnt, rev, qty in periods_sales:
        print(f"   • {label} (FY{str(yr)[2:]}): {cnt:,} records | ₹{rev:,.2f} | {qty:,} units")

    total_sales = client.command("SELECT count(*) FROM fact_sales_monthly")
    total_rev = client.command("SELECT round(sum(net_amount), 2) FROM fact_sales_monthly")
    total_qty = client.command("SELECT sum(bill_qty) FROM fact_sales_monthly")

    print(f"\nTotal Fact Sales Records: {total_sales:,}")
    print(f"Total Cumulative Revenue:   ₹{total_rev:,.2f}")
    print(f"Total Cumulative Quantity:  {total_qty:,}")

    periods_inv = client.query("""
        SELECT report_period_label, count(), sum(opening_qty), round(sum(opening_amt), 2), sum(closing_qty), round(sum(closing_amt), 2)
        FROM fact_inventory
        GROUP BY report_period_label, period_start_date
        ORDER BY period_start_date
    """).result_rows

    print("\nInventory By Period Breakdown:")
    for label, cnt, op_qty, op_amt, cl_qty, cl_amt in periods_inv:
        print(f"   • {label}: {cnt:,} records | Op Qty: {op_qty:,} (₹{op_amt:,.2f}) | Cl Qty: {cl_qty:,} (₹{cl_amt:,.2f})")

    total_inv = client.command("SELECT count(*) FROM fact_inventory")
    print(f"\nTotal Fact Inventory Records: {total_inv:,}")
    print("==================================================\n")


def main():
    start_total = time.time()
    sales_file = os.path.join("database", "data", "SALE DATA 24-26.xlsx")
    if not os.path.exists(sales_file):
        sales_file = os.path.join("data", "SALE DATA 24-26.xlsx")

    inv_file = os.path.join("database", "data", "Stock Movement DUMP Barcode wise - FY(25 & 26) - RD ( 25-09-26).xlsx")
    if not os.path.exists(inv_file):
        inv_file = os.path.join("data", "Stock Movement DUMP Barcode wise - FY(25 & 26) - RD ( 25-09-26).xlsx")

    client = get_clickhouse_client()

    # 1. Append Date Dimension
    append_dim_date(client)

    # 2. Append Sales & Products
    append_sales_and_products(client, sales_file)

    # 3. Append Inventory
    append_inventory(client, inv_file)

    # 4. Verification
    verify_clickhouse(client)

    elapsed = time.time() - start_total
    print(f"ALL FY26 DATA SUCCESSFULLY APPENDED IN {elapsed:.2f}s!")


if __name__ == "__main__":
    main()
